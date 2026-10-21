"""
Comprehensive Unit Tests for ExamSense Pipeline
================================================
Tests extractors, scanned/image rejection, limits, chunker, hybrid ranking,
not-found abstention without LLM calls, summary mode, source chips,
429 retry backoff, question quota, and key priority.
"""
import io
import os
from unittest.mock import MagicMock, patch
import fitz
from docx import Document
from ebooklib import epub
from pptx import Presentation

from core.generation.answer_generator import (
    NOT_FOUND_MESSAGE,
    generate_study_answer,
    strip_invented_citations,
)
from core.generation.llm_client import (
    check_session_question_cap,
    generate_answer,
    get_active_groq_key,
    is_user_custom_key,
)
from core.ingestion.chunker import chunk_pages
from core.ingestion.document_processor import (
    SCANNED_FILE_MESSAGE,
    UNSUPPORTED_TYPE_MESSAGE,
    process_uploaded_file,
)
from core.retrieval.hybrid_retriever import hybrid_search
from core.retrieval.vector_store import (
    _get_store,
    add_document,
    delete_document,
    get_documents,
)


# ── Helper Fixtures & Test File Generators ────────────────────────────────────

def create_text_pdf_bytes() -> bytes:
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text(
            (50, 50),
            f"Page {i+1}: Natural Language Processing and Deep Learning concepts. "
            "Transformers utilize self-attention mechanisms to process input representations in parallel. "
            "This document has substantial readable content for indexing.",
        )
    b = doc.tobytes()
    doc.close()
    return b


def create_image_only_pdf_bytes() -> bytes:
    doc = fitz.open()
    for _ in range(3):
        page = doc.new_page()
        page.insert_text((50, 50), "Hi")  # Fewer than 30 characters
    b = doc.tobytes()
    doc.close()
    return b


def create_mixed_pdf_bytes() -> bytes:
    doc = fitz.open()
    # 8 pages with plenty of text
    for i in range(8):
        page = doc.new_page()
        page.insert_text(
            (50, 50),
            f"Page {i+1}: Detailed analysis of vector embeddings and semantic search. "
            "Cosine similarity computed via dot product over normalized dense vectors.",
        )
    # 2 pages with no text (< 30 chars)
    for _ in range(2):
        doc.new_page()
    b = doc.tobytes()
    doc.close()
    return b


def create_pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Transformer Architecture and Deep Learning Foundations"
    # Table
    table_shape = slide.shapes.add_table(2, 2, 100, 100, 400, 200)
    table = table_shape.table
    table.cell(0, 0).text = "Layer Type in Neural Architecture"
    table.cell(0, 1).text = "Primary Functional Description"
    table.cell(1, 0).text = "Multi-Head Self-Attention"
    table.cell(1, 1).text = "Focuses dynamically on token representations across multiple subspaces"
    # Speaker notes
    slide.notes_slide.notes_text_frame.text = (
        "Speaker note: explain multi-head attention mechanisms, query key value projections, "
        "and how scaled dot product attention prevents gradient saturation during training."
    )

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def create_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Gradient Descent and Numerical Optimization", level=1)
    doc.add_paragraph(
        "Optimization algorithms iteratively update model weights to minimize the loss function. "
        "Stochastic gradient descent computes approximations over mini-batches, balancing stability "
        "and computational throughput across high-dimensional parameter spaces. "
        "Momentum helps accelerate gradients in vectors of consistent direction."
    )
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "Learning Rate Hyperparameter"
    tbl.rows[0].cells[1].text = "Step Size in Parameter Space"
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def create_epub_bytes() -> bytes:
    book = epub.EpubBook()
    book.set_identifier("test_epub_id")
    book.set_title("Neural Networks Primer")
    book.set_language("en")

    # Add chapter with >= 300 chars
    c1 = epub.EpubHtml(title="Backpropagation", file_name="chap1.xhtml", lang="en")
    c1.content = (
        "<h1>Backpropagation</h1>"
        "<p>The chain rule is applied to compute gradients of the loss function with respect to "
        "all layer parameters in the network. Forward propagation produces activation vectors, "
        "and backward propagation propagates error differentials back through adjacent weight layers. "
        "This algorithm is the cornerstone of modern deep learning and neural network training.</p>"
    )
    book.add_item(c1)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", c1]

    bio = io.BytesIO()
    epub.write_epub(bio, book)
    return bio.getvalue()


# ── Step 2 Tests: File Extractors & Input Validation Rules ─────────────────────

def test_extract_text_pdf():
    pdf_bytes = create_text_pdf_bytes()
    res = process_uploaded_file(pdf_bytes, "lecture_notes.pdf")
    assert res["success"] is True
    assert res["status_note"] == "Ready"
    assert len(res["chunks"]) > 0


def test_reject_image_only_pdf():
    img_pdf = create_image_only_pdf_bytes()
    res = process_uploaded_file(img_pdf, "scanned_document.pdf")
    assert res["success"] is False
    assert res["error"] == SCANNED_FILE_MESSAGE
    assert len(res["chunks"]) == 0


def test_accept_mostly_text_pdf_with_skipped_note():
    mixed_pdf = create_mixed_pdf_bytes()
    res = process_uploaded_file(mixed_pdf, "mixed_doc.pdf")
    assert res["success"] is True
    assert "2 pages/slides had no readable text and were skipped." in res["status_note"]
    assert len(res["chunks"]) > 0


def test_extract_pptx():
    pptx_bytes = create_pptx_bytes()
    res = process_uploaded_file(pptx_bytes, "slides.pptx")
    assert res["success"] is True
    full_text = " ".join(c["text"] for c in res["chunks"])
    assert "Transformer Architecture" in full_text
    assert "Attention" in full_text
    assert "Speaker note" in full_text


def test_extract_docx():
    docx_bytes = create_docx_bytes()
    res = process_uploaded_file(docx_bytes, "notes.docx")
    assert res["success"] is True
    full_text = " ".join(c["text"] for c in res["chunks"])
    assert "Gradient Descent" in full_text
    assert "Learning Rate" in full_text


def test_extract_epub():
    epub_bytes = create_epub_bytes()
    res = process_uploaded_file(epub_bytes, "book.epub")
    assert res["success"] is True
    full_text = " ".join(c["text"] for c in res["chunks"])
    assert "Backpropagation" in full_text


def test_extract_txt():
    txt_bytes = (
        b"Recurrent neural networks process sequential data one step at a time, maintaining an internal hidden state. "
        b"Long Short-Term Memory networks (LSTMs) and Gated Recurrent Units (GRUs) introduce gating mechanisms to mitigate "
        b"the vanishing gradient problem over long sequence horizons. This plain text file contains sufficient text."
    )
    res = process_uploaded_file(txt_bytes, "summary.txt")
    assert res["success"] is True
    assert len(res["chunks"]) > 0


def test_reject_unsupported_file_type():
    res = process_uploaded_file(b"\x89PNG\r\n\x1a\n", "diagram.png")
    assert res["success"] is False
    assert res["error"] == UNSUPPORTED_TYPE_MESSAGE


# ── Step 4 Tests: Chunker Metadata ────────────────────────────────────────────

def test_chunker_metadata():
    pages = [
        {"page_or_slide": 1, "text": "Section on linear regression and cost functions.", "heading": "Regression"},
        {"page_or_slide": 2, "text": "Section on logistic regression and classification.", "heading": "Classification"},
    ]
    chunks = chunk_pages(pages, doc_name="math.pdf", target_tokens=100, overlap_tokens=20)
    assert len(chunks) >= 2
    for c in chunks:
        assert c.doc_name == "math.pdf"
        assert c.page_or_slide in (1, 2)
        assert c.chunk_index is not None
        assert c.text.startswith("[")


# ── Step 3 & Step 4 Tests: In-Memory Index & Hybrid Retrieval ─────────────────

def test_in_memory_index_and_hybrid_ranking():
    # Reset store
    store = _get_store()
    store["documents"] = {}
    store["chunks"] = []

    chunks_doc1 = [
        {"chunk_id": "c1", "text": "Attention mechanisms allow transformers to weight tokens dynamically.", "metadata": {"doc_name": "ai.pdf", "page_or_slide": 1}},
        {"chunk_id": "c2", "text": "Backpropagation computes gradients through the chain rule.", "metadata": {"doc_name": "ai.pdf", "page_or_slide": 2}},
    ]
    ok, _ = add_document("ai.pdf", chunks_doc1, status_note="Ready", total_pages=2)
    assert ok is True
    assert "ai.pdf" in [d["doc_name"] for d in get_documents()]

    # Hybrid search
    results = hybrid_search(query="attention in transformers", top_k_final=2)
    assert len(results) > 0
    assert "c1" == results[0]["chunk_id"]

    # Delete document
    delete_document("ai.pdf")
    assert len(get_documents()) == 0
    empty_results = hybrid_search(query="attention in transformers", top_k_final=2)
    assert len(empty_results) == 0


# ── Step 5 Tests: Generation, Abstention, Summary Mode, and Citations ─────────

def test_not_found_makes_no_llm_call():
    # Empty store -> query should abstain without calling LLM
    store = _get_store()
    store["documents"] = {}
    store["chunks"] = []

    with patch("core.generation.answer_generator.generate_answer") as mock_llm:
        res = generate_study_answer("What is quantum entanglement?")
        assert res["not_found"] is True
        assert res["answer"] == NOT_FOUND_MESSAGE
        mock_llm.assert_not_called()


def test_citation_chips_built_from_chunks_and_invented_n_stripped():
    raw_answer = "Attention computes weighted tokens [1] and another fact [99] invented by model."
    cleaned = strip_invented_citations(raw_answer, num_chunks=1)
    assert "[1]" in cleaned
    assert "[99]" not in cleaned


def test_summary_path_skips_evidence_check():
    store = _get_store()
    store["documents"] = {}
    store["chunks"] = []

    chunks = [
        {"chunk_id": f"c_{i}", "text": f"Concept {i}: Detailed explanation of module {i}.", "metadata": {"doc_name": "course.pdf", "page_or_slide": i, "chunk_index": i}}
        for i in range(1, 10)
    ]
    add_document("course.pdf", chunks, "Ready", 9)

    with patch("core.generation.answer_generator.generate_answer") as mock_llm:
        mock_llm.return_value = {
            "answer": "This course covers modules 1 through 9 [1].",
            "error": "",
            "abstained": False,
        }
        res = generate_study_answer("Summarize this document")
        assert res["is_summary"] is True
        assert res["not_found"] is False
        mock_llm.assert_called_once()
        # Max 6 spread out chunks used in context
        called_chunks = mock_llm.call_args[1]["context_chunks"]
        assert len(called_chunks) <= 6


# ── Step 5 & Step 6 Tests: Groq 429 Retry, Quotas, and Key Priority ───────────

def test_429_retry_with_mocked_client():
    from openai import RateLimitError

    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.choices = [MagicMock()]
    mock_resp.choices[0].message.content = "Transformers use self-attention [1]."

    # First call raises RateLimitError, second succeeds
    mock_client.chat.completions.create.side_effect = [
        RateLimitError("Rate limit reached. retry-after: 0.01", response=MagicMock(status_code=429), body={}),
        mock_resp,
    ]

    with patch("core.generation.llm_client.get_active_groq_key", return_value="gsk_test_key"):
        with patch("openai.OpenAI", return_value=mock_client):
            chunks = [{"text": "Sample text", "metadata": {"doc_name": "a.pdf", "page_or_slide": 1}}]
            res = generate_answer("What is self-attention?", chunks)
            assert "Transformers use self-attention" in res["answer"]
            assert mock_client.chat.completions.create.call_count == 2


def test_per_session_cap_on_shared_key():
    import streamlit as st
    st.session_state["session_questions_count"] = 30
    st.session_state["sidebar_user_groq_key"] = ""

    can_ask, msg = check_session_question_cap()
    assert can_ask is False
    assert "reached the free session limit of 30 questions" in msg

    # With user custom key, cap is bypassed
    st.session_state["sidebar_user_groq_key"] = "gsk_user_custom_key"
    can_ask_user, _ = check_session_question_cap()
    assert can_ask_user is True
    assert is_user_custom_key() is True


def test_key_priority_order():
    import streamlit as st

    # 1. User typed key wins
    st.session_state["sidebar_user_groq_key"] = "user_key_123"
    assert get_active_groq_key() == "user_key_123"

    # 2. Falls back to environment if user key empty
    st.session_state["sidebar_user_groq_key"] = ""
    with patch.dict(os.environ, {"GROQ_API_KEY": "env_key_456"}):
        assert get_active_groq_key() == "env_key_456"
