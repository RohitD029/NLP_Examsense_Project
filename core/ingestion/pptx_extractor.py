"""
PPTX Extractor
==============
Per-slide extraction: title, text frames, recursive grouped shapes, table cells, speaker notes.
No OCR, no Tesseract.
Merges small neighbouring slides into one chunk, keeping the first slide number.
"""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_pptx(file_path: str) -> dict:
    """
    Extracts text per slide from PPTX.

    Returns dict:
      pages: list of slide dicts [{"page_or_slide": int, "text": str, "is_image_only": bool, "heading": str}]
      total_chars: int
      image_only_pages: int
      error: str | None
    """
    from pptx import Presentation
    from core.config import settings
    from core.nlp.preprocess import clean_text

    path = Path(file_path)
    try:
        prs = Presentation(str(path))
    except Exception as e:
        logger.warning(f"Cannot open PPTX {path.name}: {e}")
        return {
            "pages": [],
            "total_chars": 0,
            "image_only_pages": 0,
            "error": "This presentation appears to be corrupted or invalid.",
        }

    try:
        total_slides = len(prs.slides)
        if total_slides == 0:
            return {
                "pages": [],
                "total_chars": 0,
                "image_only_pages": 0,
                "error": "This presentation is empty.",
            }

        raw_slides = []
        image_only_count = 0
        total_chars = 0

        for slide_num, slide in enumerate(prs.slides, 1):
            parts: list[str] = []
            title = ""

            # 1. Slide Title
            try:
                if slide.shapes.title and slide.shapes.title.text:
                    title = slide.shapes.title.text.strip()
            except Exception:
                pass

            # 2. Collect shapes recursively (text frames, tables, groups)
            _collect_shapes(slide.shapes, parts)

            # 3. Speaker notes
            try:
                notes_slide = slide.notes_slide
                notes_text = notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    parts.append(f"[Notes] {notes_text}")
            except Exception:
                pass

            # Deduplicate while preserving order
            seen = set()
            clean_parts = []
            for p in parts:
                p_str = p.strip()
                if p_str and p_str not in seen:
                    seen.add(p_str)
                    clean_parts.append(p_str)

            slide_text = clean_text("\n".join(clean_parts))
            char_count = len(slide_text.strip())

            is_image_only = char_count < settings.min_pdf_page_chars
            if is_image_only:
                image_only_count += 1
            else:
                total_chars += char_count

            raw_slides.append({
                "page_or_slide": slide_num,
                "text": slide_text if not is_image_only else "",
                "is_image_only": is_image_only,
                "heading": title,
            })

        # Merge small neighbouring slides into single units, keeping the first slide number
        merged_slides = _merge_small_slides(raw_slides, target_tokens=settings.chunk_target_tokens)

        return {
            "pages": merged_slides,
            "total_chars": total_chars,
            "image_only_pages": image_only_count,
            "error": None,
        }
    except Exception as e:
        logger.error(f"Error reading PPTX {path.name}: {e}")
        return {
            "pages": [],
            "total_chars": 0,
            "image_only_pages": 0,
            "error": f"Failed to extract text from PPTX: {str(e)}",
        }


def _collect_shapes(shapes, parts: list[str]) -> None:
    """Recursively collect text from shapes, grouped shapes, and tables."""
    for shape in shapes:
        # Grouped shapes (MSO_SHAPE_TYPE.GROUP == 6)
        if getattr(shape, "shape_type", None) == 6 or hasattr(shape, "shapes"):
            try:
                _collect_shapes(shape.shapes, parts)
            except Exception:
                pass
            continue

        # Text frames
        if getattr(shape, "has_text_frame", False):
            try:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        parts.append(t)
            except Exception:
                pass

        # Tables
        if getattr(shape, "has_table", False):
            try:
                for row in shape.table.rows:
                    cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cell_texts:
                        parts.append(" | ".join(cell_texts))
            except Exception:
                pass


def _merge_small_slides(slides: list[dict], target_tokens: int = 200) -> list[dict]:
    """
    Merge consecutive slides that are individually small.
    Keeps the first slide number for the merged chunk.
    """
    if not slides:
        return []

    merged: list[dict] = []
    current_chunk = None

    for s in slides:
        if s.get("is_image_only") or not s.get("text", "").strip():
            # If skipping image-only, record as is
            if current_chunk:
                merged.append(current_chunk)
                current_chunk = None
            merged.append(s)
            continue

        text = s["text"]
        approx_tokens = max(1, len(text) // 4)

        if current_chunk is None:
            current_chunk = dict(s)
        else:
            curr_tokens = max(1, len(current_chunk["text"]) // 4)
            if curr_tokens + approx_tokens <= target_tokens:
                # Merge into current_chunk
                title_prefix = f"\n\n[{s['heading']}]\n" if s.get("heading") else "\n\n"
                current_chunk["text"] += title_prefix + text
            else:
                merged.append(current_chunk)
                current_chunk = dict(s)

    if current_chunk:
        merged.append(current_chunk)

    return merged
